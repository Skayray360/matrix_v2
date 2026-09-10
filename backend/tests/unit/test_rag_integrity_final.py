# Creado por Aldo Garcia.
"""Regresiones de fidelidad documental; dobles locales, sin inferencia real."""

from __future__ import annotations

import io
from dataclasses import replace
from types import SimpleNamespace

import pytest

from app.agents.knowledge_agent import KnowledgeAgent
from app.agents.prompts import build_answer_messages, format_evidence_block
from app.ingestion.loaders import extract_docx
from app.llm.model_policy import Intent, ModelPolicy
from app.llm.ollama_client import ChatResult
from app.memory.service import ConversationContext, ConversationTurn
from app.rag.grounding import verify_grounding
from app.rag.schemas import Evidence
from app.structured_data.tool import StructuredEvidence

pytestmark = pytest.mark.unit


def evidence(text: str, index: int = 0) -> Evidence:
    return Evidence(
        source_id=f"prestaciones/politica.md#{index}", text=text, score=0.9,
        category="prestaciones", filename="politica.md", section="Elegibilidad",
        page_or_sheet="pagina 1", document_id="doc", chunk_id=f"chunk-{index}",
    )


class ScriptedLlm:
    def __init__(self, responses: list[str]) -> None:
        self.responses = responses
        self.calls: list[dict] = []

    def chat(self, *, model: str, messages: list[dict], **kwargs) -> ChatResult:
        self.calls.append({"model": model, "messages": messages, **kwargs})
        return ChatResult(content=self.responses.pop(0), model=model, latency_ms=1)


def test_docx_tables_keep_body_order_and_own_heading():
    from docx import Document

    document = Document()
    document.add_heading("Vacaciones", level=1)
    document.add_paragraph("La antiguedad determina los dias disponibles.")
    first = document.add_table(rows=2, cols=2)
    first.cell(0, 0).text = "Antiguedad"
    first.cell(0, 1).text = "Dias"
    first.cell(1, 0).text = "5 anios"
    first.cell(1, 1).text = "20"
    document.add_paragraph("No aplica al personal temporal.")
    document.add_heading("Becas", level=1)
    second = document.add_table(rows=1, cols=1)
    second.cell(0, 0).text = "Promedio minimo: 8"
    stream = io.BytesIO()
    document.save(stream)

    result = extract_docx(stream.getvalue())
    assert [block.kind for block in result.blocks] == [
        "heading", "paragraph", "table", "paragraph", "heading", "table",
    ]
    assert result.blocks[2].section == "Vacaciones"
    assert result.blocks[2].page_or_sheet == "tabla 1"
    assert "20" in result.blocks[2].text
    assert result.blocks[5].section == "Becas"
    assert result.blocks[5].page_or_sheet == "tabla 2"


@pytest.mark.parametrize(
    ("source", "answer"),
    [
        ("El personal temporal no tiene derecho al bono.", "Tiene derecho al bono."),
        ("Si tiene cinco anios, recibe 20 dias.", "Recibe 20 dias."),
        ("El personal permanente recibe 20 dias.", "Recibe 20 dias."),
        ("Recibe 20 dias. No aplica al personal temporal.", "Recibe 20 dias."),
        ("Recibe 20 dias.\n\nNo aplica al personal temporal.", "Recibe 20 dias."),
        ("Se autoriza el bono solo con evaluacion aprobada.", "Se autoriza el bono."),
    ],
)
def test_citation_does_not_validate_missing_conditions(source: str, answer: str):
    item = evidence(source)
    report = verify_grounding(f"{answer} [[{item.source_id}]]", (item,))
    assert report.citations_valid is True
    assert report.grounded is False
    assert report.extractive_verified is False
    assert report.factual_verified is False


def test_complete_unit_preserves_cross_paragraph_exception_without_truth_claim():
    item = evidence("El personal permanente recibe 20 dias.\n\nNo aplica al personal temporal.")
    report = verify_grounding(f"{item.text} [[{item.source_id}]]", (item,))
    assert report.grounded and report.extractive_verified and report.citations_valid
    assert report.factual_verified is False


def test_valid_ids_without_a_claim_are_rejected():
    item = evidence("La prima es del 25%.")
    assert not verify_grounding(f"[[{item.source_id}]]", (item,)).grounded


def test_same_source_id_cannot_disguise_different_document():
    first = evidence("No se autoriza el bono.")
    second = replace(first, text="Se autoriza el bono.", document_id="another")
    report = verify_grounding(f"{second.text} [[{second.source_id}]]", (first, second))
    assert not report.grounded
    assert "ambiguo" in report.reason


def test_sql_keeps_full_row_and_column_meaning():
    result = StructuredEvidence(
        source_id="db:rh/permisos", source="rh", entity="permisos",
        columns=("tipo", "dias"), rows=(("sin goce", 12),), row_count=1, truncated=False,
    )
    assert not verify_grounding("12 [[db:rh/permisos]]", (), structured=(result,)).grounded
    assert not verify_grounding("dias: 12 [[db:rh/permisos]]", (), structured=(result,)).grounded
    valid = verify_grounding("tipo: sin goce | dias: 12 [[db:rh/permisos]]", (), structured=(result,))
    assert valid.extractive_verified and not valid.factual_verified


def test_normal_prompt_keeps_exception_after_old_2200_cut():
    item = evidence("El procedimiento exige comprobantes. " * 80 + "EXCEPCION: no aplica al personal temporal.")
    messages = build_answer_messages(question="A quien aplica?", evidences=(item,))
    assert item.text in messages[-1]["content"]
    assert "EXCEPCION: no aplica al personal temporal." in messages[-1]["content"]
    with pytest.raises(ValueError, match="evidencia completa excede"):
        format_evidence_block((item,), max_chars_each=2200)


def test_documental_prompt_budget_omits_whole_unit_and_reports_selection():
    class BudgetPolicy(ModelPolicy):
        def generation_profile(self, *_args, **_kwargs):
            return SimpleNamespace(num_ctx=6000, max_tokens=1000, temperature=0.1)

    first = evidence("Regla A completa. " * 260 + "No aplica al personal temporal.")
    second = evidence("Regla B completa. " * 260 + "Solo con autorizacion escrita.", 1)
    llm = ScriptedLlm([f"{first.text} [[{first.source_id}]]"])
    agent = KnowledgeAgent(llm=llm, policy=BudgetPolicy())
    result = agent.synthesize(question="Que reglas aplican?", evidences=(first, second), model_name="gemma4:latest")
    messages = llm.calls[0]["messages"]
    assert first.text in messages[-1]["content"]
    assert second.source_id not in messages[-1]["content"]
    assert sum(len(message["content"]) for message in messages) <= agent._input_budget_chars("gemma4:latest", Intent.DOCUMENTAL)
    assert "seleccion de unidades completas" in result.answer
    assert result.grounding.extractive_verified


def test_oversized_documental_unit_does_not_reach_runtime():
    item = evidence("Regla muy extensa. " * 10000 + "No aplica al personal temporal.")
    llm = ScriptedLlm([])
    result = KnowledgeAgent(llm=llm, policy=ModelPolicy()).synthesize(
        question="A quien aplica?", evidences=(item,), model_name="gemma4:latest",
    )
    assert not llm.calls
    assert not result.grounding.grounded
    assert "presupuesto" in result.answer


def test_extractive_fallback_keeps_exception_after_old_420_cut():
    item = evidence("El procedimiento exige comprobantes. " * 30 + "No aplica al personal temporal.")
    llm = ScriptedLlm(["Resumen sin fuentes.", "Tampoco tiene fuentes."])
    result = KnowledgeAgent(llm=llm, policy=ModelPolicy()).synthesize(
        question="Resume el documento", evidences=(item,), model_name="gemma4:latest",
        intent=Intent.DOCUMENT_SUMMARY,
    )
    assert item.text in result.answer
    assert "No aplica al personal temporal." in result.answer
    assert result.grounding.extractive_verified
    assert not result.grounding.factual_verified
    assert len(llm.calls) == 2


def test_rejected_documental_paraphrase_keeps_original_abstention_contract():
    item = evidence("Si tiene cinco anios, recibe 20 dias. No aplica al personal temporal.")
    llm = ScriptedLlm([f"Recibe 20 dias [[{item.source_id}]]."] * 2)
    result = KnowledgeAgent(llm=llm, policy=ModelPolicy()).synthesize(
        question="Cuantos dias?", evidences=(item,), model_name="gemma4:latest",
    )
    assert item.text not in result.answer
    assert "No cuento con informacion documental suficiente" in result.answer
    assert not result.grounding.grounded
    assert not result.grounding.extractive_verified
    assert not result.grounding.factual_verified
    assert result.cited_source_ids == ()
    assert len(llm.calls) == 2


@pytest.mark.parametrize(("source", "answer"), [("- 12", "12"), ("-12", "12"), (".5", "5"), ("+12", "12")])
def test_normalization_preserves_numeric_signs_and_initial_decimal(source: str, answer: str):
    item = evidence(source)
    assert not verify_grounding(f"{answer} [[{item.source_id}]]", (item,)).grounded
    assert verify_grounding(f"{source} [[{item.source_id}]]", (item,)).extractive_verified


def test_complete_evidence_has_priority_over_long_memory():
    item = evidence("La politica exige autorizacion escrita. " * 90)
    memory = ConversationContext(
        conversation_id="conversation",
        summary="Resumen historico. " * 290,
        turns=tuple(ConversationTurn(role="user", content="Dialogo previo. " * 80) for _ in range(6)),
    )
    llm = ScriptedLlm([f"{item.text} [[{item.source_id}]]"])
    result = KnowledgeAgent(llm=llm, policy=ModelPolicy()).synthesize(
        question="Que exige la politica?", evidences=(item,), memory=memory, model_name="gemma4:latest",
    )
    assert len(llm.calls) == 1
    prompt = llm.calls[0]["messages"][-1]["content"]
    assert item.text in prompt
    assert "MEMORIA_CONVERSACION" not in prompt
    assert result.grounding.extractive_verified


def test_failed_documental_answer_does_not_promote_irrelevant_retrieved_source():
    item = evidence("El comedor abre de lunes a viernes.")
    llm = ScriptedLlm(["Respuesta sin citas."] * 2)
    result = KnowledgeAgent(llm=llm, policy=ModelPolicy()).synthesize(
        question="Cuantos dias de permiso de maternidad corresponden?",
        evidences=(item,), model_name="gemma4:latest",
    )
    assert "No cuento con informacion documental suficiente" in result.answer
    assert "comedor" not in result.answer
    assert not result.grounding.grounded
    assert not result.cited_source_ids
